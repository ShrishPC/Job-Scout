import os
import tempfile
import fitz
import docx
import pytest
from unittest.mock import patch
from app.services.resume_service import parse_resume_to_markdown

def test_parse_resume_txt_and_md():
    with tempfile.NamedTemporaryFile("w+", suffix=".txt", delete=False) as f:
        f.write("Candidate: John Doe\nSkills: Python, SQL")
        txt_path = f.name

    with tempfile.NamedTemporaryFile("w+", suffix=".md", delete=False) as f:
        f.write("# Jane Doe\n- Expertise: Docker, Kubernetes")
        md_path = f.name

    try:
        txt_out = parse_resume_to_markdown(txt_path)
        assert "John Doe" in txt_out
        assert "Skills: Python, SQL" in txt_out

        md_out = parse_resume_to_markdown(md_path)
        assert "Jane Doe" in md_out
        assert "Docker, Kubernetes" in md_out
    finally:
        os.remove(txt_path)
        os.remove(md_path)

def test_parse_resume_pdf():
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
        pdf_path = f.name

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), "PDF Candidate Resume\nExperience: 5 years in AWS cloud.")
    doc.save(pdf_path)
    doc.close()

    try:
        pdf_out = parse_resume_to_markdown(pdf_path)
        assert "PDF Candidate Resume" in pdf_out
        assert "AWS cloud" in pdf_out
    finally:
        os.remove(pdf_path)

def test_parse_resume_pdf_fallback_extractor():
    # Force MarkItDown to fail to test PyMuPDF secondary fallback
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
        pdf_path = f.name

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), "Fallback Extracted Content via PyMuPDF")
    doc.save(pdf_path)
    doc.close()

    try:
        with patch("app.services.resume_service.MarkItDown") as mock_md:
            mock_md.side_effect = RuntimeError("MarkItDown failed")
            out = parse_resume_to_markdown(pdf_path)
            assert "Fallback Extracted Content via PyMuPDF" in out
    finally:
        os.remove(pdf_path)

def test_parse_resume_docx():
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        docx_path = f.name

    doc = docx.Document()
    doc.add_heading("Docx Candidate Resume", 0)
    doc.add_paragraph("Skills: React, TypeScript, Next.js")
    doc.save(docx_path)

    try:
        docx_out = parse_resume_to_markdown(docx_path)
        assert "Docx Candidate Resume" in docx_out
        assert "React, TypeScript" in docx_out
    finally:
        os.remove(docx_path)

def test_parse_resume_docx_fallback_extractor():
    # Force MarkItDown to fail to test python-docx secondary fallback
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        docx_path = f.name

    doc = docx.Document()
    doc.add_heading("Fallback Docx Candidate", 0)
    doc.add_paragraph("Fallback python-docx skills: Golang, Rust")
    doc.save(docx_path)

    try:
        with patch("app.services.resume_service.MarkItDown") as mock_md:
            mock_md.side_effect = RuntimeError("MarkItDown failed")
            out = parse_resume_to_markdown(docx_path)
            assert "Fallback python-docx skills" in out
    finally:
        os.remove(docx_path)

def test_parse_resume_raw_fallback():
    # File with unknown extension that falls back to raw binary decode
    with tempfile.NamedTemporaryFile("w+", suffix=".dat", delete=False) as f:
        f.write("Raw decoded binary resume text content.")
        dat_path = f.name

    try:
        with patch("app.services.resume_service.MarkItDown") as mock_md:
            mock_md.side_effect = RuntimeError("MarkItDown failed")
            out = parse_resume_to_markdown(dat_path)
            assert "Raw decoded binary resume text" in out
    finally:
        os.remove(dat_path)

def test_parse_resume_nonexistent_file():
    with pytest.raises(FileNotFoundError):
        parse_resume_to_markdown("/nonexistent/path/to/resume.pdf")

import sys
import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import subprocess

prs = Presentation()
prs.slide_width = Inches(13.333)  # 16:9 widescreen
prs.slide_height = Inches(7.5)

# Lively, Modern Academic Tech Palette
BG_CANVAS = RGBColor(248, 250, 252)       # #F8FAFC Crisp light slate
CARD_WHITE = RGBColor(255, 255, 255)      # Pure White Card
CARD_BORDER = RGBColor(226, 232, 240)     # #E2E8F0 Subtle Border Gray
TEXT_DARK = RGBColor(15, 23, 42)          # #0F172A Slate-900 (High contrast)
TEXT_BODY = RGBColor(51, 65, 85)          # #334155 Slate-700
TEXT_MUTED = RGBColor(100, 116, 139)      # #64748B Slate-500

# Vibrant Accent Palette
COLOR_INDIGO = RGBColor(37, 99, 235)      # #2563EB Royal Cobalt / Primary Accent
COLOR_INDIGO_BG = RGBColor(239, 246, 255) # #EFF6FF Tint
COLOR_PURPLE = RGBColor(124, 58, 237)     # #7C3AED Electric Violet / AI
COLOR_PURPLE_BG = RGBColor(245, 243, 255) # #F5F3FF Tint
COLOR_EMERALD = RGBColor(16, 185, 129)    # #10B981 Vivid Emerald / Metrics
COLOR_EMERALD_BG = RGBColor(236, 253, 245)# #ECFDF5 Tint
COLOR_CORAL = RGBColor(234, 88, 12)       # #EA580C Sunset Coral / Problem
COLOR_CORAL_BG = RGBColor(255, 247, 237)  # #FFF7ED Tint
COLOR_CYAN = RGBColor(2, 132, 199)        # #0284C7 Sky Cyan / Pipelines
COLOR_CYAN_BG = RGBColor(240, 249, 255)   # #F0F9FF Tint

def apply_lively_background(slide):
    # Base crisp background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_CANVAS
    bg.line.fill.background()
    
    # Top dynamic gradient-feel accent stripe
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = COLOR_INDIGO
    top_bar.line.fill.background()
    return bg

def add_lively_header(slide, title_text, category_pill="JOB SCOUT // THESIS DEFENSE", pill_color=COLOR_INDIGO, pill_bg=COLOR_INDIGO_BG):
    # Pill Tag
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.35), Inches(3.2), Inches(0.35))
    pill.fill.solid()
    pill.fill.fore_color.rgb = pill_bg
    pill.line.color.rgb = pill_color
    pill.line.width = Pt(1.0)
    p_pill = pill.text_frame.paragraphs[0]
    p_pill.alignment = PP_ALIGN.CENTER
    p_pill.text = category_pill.upper()
    p_pill.font.name = "Arial"
    p_pill.font.size = Pt(9.5)
    p_pill.font.bold = True
    p_pill.font.color.rgb = pill_color

    # Title
    t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.7))
    tf_t = t_box.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text
    p_t.font.name = "Arial"
    p_t.font.size = Pt(23)
    p_t.font.bold = True
    p_t.font.color.rgb = TEXT_DARK

def create_lively_card(slide, left, top, width, height, title=None, accent_color=COLOR_INDIGO, bg_color=CARD_WHITE):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = CARD_BORDER
    card.line.width = Pt(1.0)
    
    # Accent top border strip
    if accent_color:
        strip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, Inches(0.08))
        strip.fill.solid()
        strip.fill.fore_color.rgb = accent_color
        strip.line.fill.background()
    
    if title:
        tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), width - Inches(0.4), Inches(0.4))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.text = title.upper()
        p.font.name = "Arial"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = accent_color if accent_color else TEXT_DARK
    return card

blank_slide_layout = prs.slide_layouts[6]

# =============================================================
# SLIDE 1: TITLE SLIDE (LIVELY EXECUTIVE LAUNCH)
# =============================================================
s1 = prs.slides.add_slide(blank_slide_layout)
apply_lively_background(s1)

# Main Hero Card
hero = create_lively_card(s1, Inches(1.0), Inches(1.2), Inches(11.333), Inches(5.3), accent_color=COLOR_INDIGO)

# Category Banner
badge = s1.shapes.add_textbox(Inches(1.5), Inches(1.6), Inches(10.333), Inches(0.4))
p_b = badge.text_frame.paragraphs[0]
p_b.text = "MASTER OF SCIENCE IN APPLIED DATA SCIENCE  |  FINAL PROJECT DEFENSE"
p_b.font.name = "Arial"
p_b.font.size = Pt(11)
p_b.font.bold = True
p_b.font.color.rgb = COLOR_INDIGO

# Hero Title
m_title = s1.shapes.add_textbox(Inches(1.5), Inches(2.0), Inches(10.333), Inches(1.6))
p_m = m_title.text_frame.paragraphs[0]
p_m.text = "JOB SCOUT"
p_m.font.name = "Arial"
p_m.font.size = Pt(46)
p_m.font.bold = True
p_m.font.color.rgb = TEXT_DARK

p_sub = m_title.text_frame.add_paragraph()
p_sub.text = "AI-Powered Job Discovery & ATS Optimization Platform"
p_sub.font.name = "Arial"
p_sub.font.size = Pt(21)
p_sub.font.bold = True
p_sub.font.color.rgb = COLOR_INDIGO

# 3 Lively Highlight Badges
badge_data = [
    ("126 Automated Tests", "100% Passing (14.6s CI)", COLOR_EMERALD, COLOR_EMERALD_BG, Inches(1.5)),
    ("Sub-10ms Hybrid Search", "9.08 ms on 100k Rows", COLOR_INDIGO, COLOR_INDIGO_BG, Inches(5.0)),
    ("100% Air-Gapped Local AI", "Zero Cloud API Fees", COLOR_PURPLE, COLOR_PURPLE_BG, Inches(8.5))
]
for b_top, b_sub, c_fg, c_bg, b_left in badge_data:
    b_card = slide_badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, b_left, Inches(3.8), Inches(3.2), Inches(0.8))
    slide_badge.fill.solid()
    slide_badge.fill.fore_color.rgb = c_bg
    slide_badge.line.color.rgb = c_fg
    slide_badge.line.width = Pt(1.0)
    tf = slide_badge.text_frame
    p1 = tf.paragraphs[0]
    p1.alignment = PP_ALIGN.CENTER
    p1.text = b_top
    p1.font.name = "Arial"
    p1.font.size = Pt(11)
    p1.font.bold = True
    p1.font.color.rgb = c_fg
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.text = b_sub
    p2.font.name = "Calibri"
    p2.font.size = Pt(9.5)
    p2.font.color.rgb = TEXT_BODY

# Author & Repo Footer Box
footer = s1.shapes.add_textbox(Inches(1.5), Inches(4.9), Inches(10.333), Inches(1.3))
tf_foot = footer.text_frame
pf1 = tf_foot.paragraphs[0]
pf1.text = "Candidate: Shrish Prasad Chakraborty (Reg. No: RA2532014010136)   |   Supervisor: Dr. R. Thilagavathy"
pf1.font.name = "Arial"
pf1.font.size = Pt(12)
pf1.font.bold = True
pf1.font.color.rgb = TEXT_DARK

pf2 = tf_foot.add_paragraph()
pf2.text = "Department of Computer Applications   |   SRM Institute of Science & Technology, Chennai"
pf2.font.name = "Calibri"
pf2.font.size = Pt(11)
pf2.font.color.rgb = TEXT_MUTED

pf3 = tf_foot.add_paragraph()
pf3.text = "Codebase: https://github.com/ShrishPC/Job-Scout.git   |   Host Execution Target: /home/rishav/job-scout/"
pf3.font.name = "Calibri"
pf3.font.size = Pt(10)
pf3.font.color.rgb = COLOR_INDIGO

# =============================================================
# SLIDE 2: PROBLEM STATEMENT & OBJECTIVES
# =============================================================
s2 = prs.slides.add_slide(blank_slide_layout)
apply_lively_background(s2)
add_lively_header(s2, "The Real-World Crisis & Core Engineering Objectives", category_pill="PROBLEM VS SOLUTION", pill_color=COLOR_CORAL, pill_bg=COLOR_CORAL_BG)

# Left Card: Problem (Sunset Coral Warmth)
create_lively_card(s2, Inches(0.8), Inches(1.5), Inches(5.7), Inches(5.4), title="THE RECRUITMENT ASYMMETRY CRISIS", accent_color=COLOR_CORAL, bg_color=CARD_WHITE)
tb_l = s2.shapes.add_textbox(Inches(1.0), Inches(2.1), Inches(5.3), Inches(4.6))
tf_l = tb_l.text_frame
tf_l.word_wrap = True

probs = [
    ("75% Unforgiving ATS Rejection Rate", "Brittle keyword filters discard highly qualified talent over minor lexical phrasing differences (e.g., 'Go' vs 'Golang') without evaluating semantic competence."),
    ("Severe Privacy & PII Exposure Risks", "Commercial cloud resume tools upload unredacted work histories, contacts, and confidential code to third-party cloud LLM APIs without independent security guarantees."),
    ("SaaS Paywalls & Portal Fragmentation", "Incumbent tools lock essential optimization behind $20-$50/month paywalls, while job seekers face severe fatigue juggling multiple disconnected job boards.")
]
for i, (head, desc) in enumerate(probs):
    p = tf_l.paragraphs[0] if i == 0 else tf_l.add_paragraph()
    p.space_after = Pt(12)
    r1 = p.add_run()
    r1.text = "⚠️ " + head + "\n"
    r1.font.name = "Arial"
    r1.font.size = Pt(11.5)
    r1.font.bold = True
    r1.font.color.rgb = COLOR_CORAL
    r2 = p.add_run()
    r2.text = desc
    r2.font.name = "Calibri"
    r2.font.size = Pt(10.5)
    r2.font.color.rgb = TEXT_BODY

# Right Card: Objectives (Emerald / Mint Warmth)
create_lively_card(s2, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.4), title="ENGINEERED SYSTEM OBJECTIVES", accent_color=COLOR_EMERALD, bg_color=CARD_WHITE)
tb_r = s2.shapes.add_textbox(Inches(7.0), Inches(2.1), Inches(5.3), Inches(4.6))
tf_r = tb_r.text_frame
tf_r.word_wrap = True

objs = [
    ("Autonomous Multi-Source Harvesting", "Engineer a resilient Celery/Playwright task pipeline across 5 portals (LinkedIn, Indeed, Naukri, RemoteOK, WWR) with automatic deduplication."),
    ("Sub-10ms Two-Stage Hybrid Matchmaking", "Fuse 70% dense semantic embeddings (all-MiniLM-L6-v2 via pgvector HNSW) with 30% keyword frequency for lightning-fast candidate ranking."),
    ("Deterministic 100-Point ATS Rubric", "Deliver an explainable, mathematical scorecard across Hard/Soft Skills, Experience, Education, and Formatting with zero AI hallucination."),
    ("Air-Gapped On-Device Generative Tailoring", "Execute local RAG resume synthesis via Llama-3.2-3B with real-time SSE streaming, ensuring 100% data sovereignty and zero API costs.")
]
for i, (head, desc) in enumerate(objs):
    p = tf_r.paragraphs[0] if i == 0 else tf_r.add_paragraph()
    p.space_after = Pt(10)
    r1 = p.add_run()
    r1.text = "✅ " + head + "\n"
    r1.font.name = "Arial"
    r1.font.size = Pt(11)
    r1.font.bold = True
    r1.font.color.rgb = COLOR_EMERALD
    r2 = p.add_run()
    r2.text = desc
    r2.font.size = Pt(10)
    r2.font.color.rgb = TEXT_BODY

# =============================================================
# SLIDE 3: EXISTING SYSTEM VS PROPOSED SYSTEM
# =============================================================
s3 = prs.slides.add_slide(blank_slide_layout)
apply_lively_background(s3)
add_lively_header(s3, "Competitive Analysis: Commercial Tools vs. Job Scout", category_pill="MARKET BENCHMARK", pill_color=COLOR_INDIGO, pill_bg=COLOR_INDIGO_BG)

headers = ["Operational Dimension", "Commercial ATS Tools (Jobscan / Teal / Rezi)", "Job Scout Platform (Proposed Architecture)"]
rows = [
    ["Search & Matching", "Exact keyword matching; misses conceptual synonyms", "Two-Stage Hybrid: 70% Semantic (pgvector HNSW) + 30% Keyword"],
    ["Candidate Privacy", "Uploaded to third-party cloud servers; unverified PII storage", "100% Air-gapped on-device execution; zero external network egress"],
    ["Cost & Accessibility", "Recurring SaaS paywalls ($20 to $50 / month per applicant)", "Zero paid API usage; fully self-hosted open-source software"],
    ["ATS Diagnostics", "Opaque proprietary percentage score; non-actionable suggestions", "Deterministic 100-point rubric; exact keyword gap analysis"],
    ["Content Generation", "Commercial cloud LLMs (OpenAI API) or static templates", "On-device local RAG via Llama-3.2-3B with live SSE streaming"],
    ["Job Harvesting", "Manual browsing or locked to proprietary platform partners", "Autonomous Celery task queue harvesting 5 live job portals"]
]

tbl_shape = s3.shapes.add_table(len(rows) + 1, 3, Inches(0.8), Inches(1.6), Inches(11.733), Inches(5.2))
tbl = tbl_shape.table
tbl.columns[0].width = Inches(2.333)
tbl.columns[1].width = Inches(4.7)
tbl.columns[2].width = Inches(4.7)

for col_idx, h in enumerate(headers):
    cell = tbl.cell(0, col_idx)
    cell.fill.solid()
    cell.fill.fore_color.rgb = COLOR_INDIGO if col_idx == 2 else TEXT_DARK
    p = cell.text_frame.paragraphs[0]
    p.text = h
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)

for r_idx, row in enumerate(rows):
    bg_c = CARD_WHITE if r_idx % 2 == 0 else RGBColor(241, 245, 249)
    for c_idx, val in enumerate(row):
        cell = tbl.cell(r_idx + 1, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = bg_c
        p = cell.text_frame.paragraphs[0]
        p.text = val
        p.font.name = "Calibri"
        p.font.size = Pt(10)
        if c_idx == 2:
            p.font.bold = True
            p.font.color.rgb = COLOR_INDIGO
        elif c_idx == 1 and r_idx < 3:
            p.font.color.rgb = COLOR_CORAL
        else:
            p.font.color.rgb = TEXT_BODY

# =============================================================
# SLIDE 4: SYSTEM ARCHITECTURE (WITH REAL DIAGRAM)
# =============================================================
s4 = prs.slides.add_slide(blank_slide_layout)
apply_lively_background(s4)
add_lively_header(s4, "System Architecture: Decoupled 4-Tier Microservice Topology", category_pill="SYSTEM TOPOLOGY", pill_color=COLOR_CYAN, pill_bg=COLOR_CYAN_BG)

# Left Column: 4 Clean Tier Cards
tier_info = [
    ("Tier 1: Presentation (Next.js 14)", "Responsive SPA, Server-Sent Events (SSE) token listener, dynamic Kanban board.", COLOR_INDIGO, Inches(1.5)),
    ("Tier 2: Application API (FastAPI)", "Asynchronous ASGI service, Pydantic validation schemas, two-stage SQL engine.", COLOR_CYAN, Inches(2.8)),
    ("Tier 3: Vector Database (PostgreSQL 16)", "pgvector HNSW index (m=16, ef=128), session ef_search=400 clamping, relational store.", COLOR_PURPLE, Inches(4.1)),
    ("Tier 4: Distributed Workers (Redis & Celery)", "Headless Playwright scraper daemon, task broker, SHA-256 LLM response caching.", COLOR_EMERALD, Inches(5.4))
]
for t_title, t_desc, c_color, top_pos in tier_info:
    card = create_lively_card(s4, Inches(0.8), top_pos, Inches(5.5), Inches(1.15), title=t_title, accent_color=c_color)
    tb = s4.shapes.add_textbox(Inches(1.0), top_pos + Inches(0.4), Inches(5.1), Inches(0.65))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = t_desc
    p.font.name = "Calibri"
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_BODY

# Right Column: High-Res Architecture Diagram
arch_card = create_lively_card(s4, Inches(6.6), Inches(1.5), Inches(5.9), Inches(5.4), title="VERIFIED SYSTEM DATA FLOW", accent_color=COLOR_INDIGO)
img_path = "/home/rishav/job-scout/assets/system_architecture.jpg"
if os.path.exists(img_path):
    s4.shapes.add_picture(img_path, Inches(6.8), Inches(2.1), Inches(5.5), Inches(4.5))

# =============================================================
# SLIDE 5: TECHNOLOGIES USED
# =============================================================
s5 = prs.slides.add_slide(blank_slide_layout)
apply_lively_background(s5)
add_lively_header(s5, "Technology Stack: Purpose-Built On-Device Engineering", category_pill="TECH STACK & TOOLS", pill_color=COLOR_PURPLE, pill_bg=COLOR_PURPLE_BG)

tech_quads = [
    ("CORE BACKEND & DATABASE", [
        ("FastAPI & Uvicorn:", "High-concurrency async Python framework with automatic OpenAPI documentation."),
        ("SQLAlchemy 2.0 ORM:", "Optimized connection pool (pool_size=10, max_overflow=20) with pre-ping validation."),
        ("PostgreSQL 16 + pgvector:", "Hybrid store combining relational job data with 384-dimensional vector indexing.")
    ], COLOR_INDIGO, Inches(0.8), Inches(1.5)),
    ("EMBEDDINGS & LOCAL GENERATIVE AI", [
        ("all-MiniLM-L6-v2:", "384-dimensional sentence transformer generating dense semantic vectors in ~5ms."),
        ("Llama-3.2-3B-Instruct:", "SOTA small language model running 100% locally in bfloat16 via PyTorch."),
        ("TinyLlama-1.1B:", "Secondary automated fallback model ensuring graceful degradation on low-RAM hosts.")
    ], COLOR_PURPLE, Inches(6.8), Inches(1.5)),
    ("ASYNC WORKERS & HARVESTING", [
        ("Redis 7.2 Message Broker:", "In-memory task queue broker and cryptographic SHA-256 generation cache."),
        ("Celery 5.3 Distributed Queue:", "Non-blocking background task runner scheduling concurrent scraper jobs."),
        ("Playwright 1.44 Headless:", "Browser automation overcoming dynamic client-rendered job board interfaces.")
    ], COLOR_CYAN, Inches(0.8), Inches(4.3)),
    ("FRONTEND, DEVOPS & TESTING", [
        ("Next.js 14 & Tailwind CSS:", "Component-driven responsive interface styled with Lucide icons and SSE hooks."),
        ("Docker & Docker Compose:", "Multi-container orchestration providing reproducible, zero-friction deployments."),
        ("pytest Test Framework:", "126 automated test cases establishing 79% verified backend line coverage.")
    ], COLOR_EMERALD, Inches(6.8), Inches(4.3))
]

for title, items, c_color, left, top in tech_quads:
    create_lively_card(s5, left, top, Inches(5.7), Inches(2.6), title=title, accent_color=c_color)
    tb = s5.shapes.add_textbox(left + Inches(0.2), top + Inches(0.55), Inches(5.3), Inches(1.9))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, (b_title, b_desc) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(4)
        r1 = p.add_run()
        r1.text = "• " + b_title + " "
        r1.font.name = "Arial"
        r1.font.bold = True
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = c_color
        r2 = p.add_run()
        r2.text = b_desc
        r2.font.name = "Calibri"
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = TEXT_BODY

# =============================================================
# SLIDE 6: IMPLEMENTATION & KEY FEATURES (WITH UI EMBED)
# =============================================================
s6 = prs.slides.add_slide(blank_slide_layout)
apply_lively_background(s6)
add_lively_header(s6, "Implementation Highlights: 4 Production Modules", category_pill="CORE MODULES", pill_color=COLOR_INDIGO, pill_bg=COLOR_INDIGO_BG)

# Left Column: Module Details
mod_cards = [
    ("01 // THE VAULT (Resume Ingestion)", "Multi-engine parser (PDF/DOCX) stripping formatting artifacts into clean Markdown; generates 384-d vectors via SentenceTransformers.", COLOR_INDIGO, Inches(1.5)),
    ("02 // THE HUNT (Two-Stage Matchmaking)", "Stage 1 retrieves top 200 candidates via HNSW in 2.88 ms; Stage 2 scores keywords in-memory (9.08 ms total API latency).", COLOR_CYAN, Inches(2.8)),
    ("03 // THE RADAR (100-Point ATS Rubric)", "Explainable rubric: Hard Skills (40), Soft (20), Exp (25), Edu (10), Format (5) with frequency-weighted keyword gap analysis.", COLOR_EMERALD, Inches(4.1)),
    ("04 // THE COPILOT (Local RAG Tailoring)", "Llama-3.2-3B synthesizes tailored executive summary + 3 STAR bullets streamed live via SSE with zero cloud API keys.", COLOR_PURPLE, Inches(5.4))
]
for m_title, m_desc, c_color, top_pos in mod_cards:
    card = create_lively_card(s6, Inches(0.8), top_pos, Inches(5.5), Inches(1.15), title=m_title, accent_color=c_color)
    tb = s6.shapes.add_textbox(Inches(1.0), top_pos + Inches(0.4), Inches(5.1), Inches(0.65))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = m_desc
    p.font.name = "Calibri"
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_BODY

# Right Column: Real UI Screenshot
ui_card = create_lively_card(s6, Inches(6.6), Inches(1.5), Inches(5.9), Inches(5.4), title="LIVE RADAR & ATS SCORECARD UI", accent_color=COLOR_EMERALD)
img_radar = "/home/rishav/job-scout/assets/radar.png"
if os.path.exists(img_radar):
    s6.shapes.add_picture(img_radar, Inches(6.8), Inches(2.1), Inches(5.5), Inches(4.5))

# =============================================================
# SLIDE 7: RESULTS & PERFORMANCE EVALUATION (VIBRANT METRICS)
# =============================================================
s7 = prs.slides.add_slide(blank_slide_layout)
apply_lively_background(s7)
add_lively_header(s7, "Empirical Validation: Measured Benchmarks & Metrics", category_pill="EMPIRICAL RESULTS", pill_color=COLOR_EMERALD, pill_bg=COLOR_EMERALD_BG)

# 4 High-Impact Metric Badges
metric_data = [
    ("9.08 ms", "API Match Latency", "POST /jobs/matches on 100k rows (98.1x speedup)", COLOR_EMERALD, COLOR_EMERALD_BG, Inches(0.8)),
    ("0.9565", "HNSW Recall@10", "Tuned at ef_search=400 (vs 0.58 default drop)", COLOR_INDIGO, COLOR_INDIGO_BG, Inches(3.8)),
    ("0.985", "Retrieval NDCG@10", "28x560 Benchmark (P@5 = 0.986, MRR = 1.0)", COLOR_CYAN, COLOR_CYAN_BG, Inches(6.8)),
    ("30,637x", "Redis Cache Acceleration", "61.33 s Cold CPU down to 2.00 ms Warm Hit", COLOR_PURPLE, COLOR_PURPLE_BG, Inches(9.8))
]
for val, title, desc, c_fg, c_bg, left_pos in metric_data:
    card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, Inches(1.5), Inches(2.7), Inches(1.6))
    card.fill.solid()
    card.fill.fore_color.rgb = c_bg
    card.line.color.rgb = c_fg
    card.line.width = Pt(1.2)
    tf = card.text_frame
    p1 = tf.paragraphs[0]
    p1.alignment = PP_ALIGN.CENTER
    p1.text = val
    p1.font.name = "Arial"
    p1.font.size = Pt(28)
    p1.font.bold = True
    p1.font.color.rgb = c_fg
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.text = title
    p2.font.name = "Arial"
    p2.font.size = Pt(10)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_DARK
    p3 = tf.add_paragraph()
    p3.alignment = PP_ALIGN.CENTER
    p3.text = desc
    p3.font.name = "Calibri"
    p3.font.size = Pt(8.5)
    p3.font.color.rgb = TEXT_MUTED

# Bottom Row: Test Suite & Grounding Cards
create_lively_card(s7, Inches(0.8), Inches(3.4), Inches(5.7), Inches(3.5), title="TEST SUITE & SECURITY AUDIT", accent_color=COLOR_INDIGO)
tb_ts = s7.shapes.add_textbox(Inches(1.0), Inches(3.9), Inches(5.3), Inches(2.8))
tf_ts = tb_ts.text_frame
tf_ts.word_wrap = True
ts_lines = [
    ("126 Automated Tests Passing:", " 100% pass rate in 14.60 seconds via pytest CI."),
    ("79% Backend Line Coverage:", " 100% on models & config; 88% on ATS; 84% on parser."),
    ("Air-Gapped Offline Execution:", " Full pipeline finishes in 145.55 ms with network blocked."),
    ("Hardened Security Layer:", " Parameterized SQL bindings, admin API key, 10MB upload limit.")
]
for i, (b_title, b_desc) in enumerate(ts_lines):
    p = tf_ts.paragraphs[0] if i == 0 else tf_ts.add_paragraph()
    p.space_after = Pt(6)
    r1 = p.add_run()
    r1.text = "• " + b_title + " "
    r1.font.bold = True
    r1.font.size = Pt(10)
    r1.font.color.rgb = COLOR_INDIGO
    r2 = p.add_run()
    r2.text = b_desc
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = TEXT_BODY

create_lively_card(s7, Inches(6.8), Inches(3.4), Inches(5.7), Inches(3.5), title="FACTUALITY & HUMAN REVIEW", accent_color=COLOR_PURPLE)
tb_gw = s7.shapes.add_textbox(Inches(7.0), Inches(3.9), Inches(5.3), Inches(2.8))
tf_gw = tb_gw.text_frame
tf_gw.word_wrap = True
gw_lines = [
    ("9.5% Unsupported Claims:", " Only 8/84 claims unsupported for Llama-3.2 vs 29.8% for TinyLlama."),
    ("4.20 / 5.00 Blind Human Rating:", " Evaluated across 20 samples for relevance, tone, and fit (vs 2.75)."),
    ("CPU Throughput Constraint:", " 1.82 tok/s CPU throughput bound by 6.42 GB bfloat16 RAM bandwidth."),
    ("Zero Third-Party Data Leakage:", " Fully local architecture eliminates recurring API costs and PII risk.")
]
for i, (b_title, b_desc) in enumerate(gw_lines):
    p = tf_gw.paragraphs[0] if i == 0 else tf_gw.add_paragraph()
    p.space_after = Pt(6)
    r1 = p.add_run()
    r1.text = "• " + b_title + " "
    r1.font.bold = True
    r1.font.size = Pt(10)
    r1.font.color.rgb = COLOR_PURPLE
    r2 = p.add_run()
    r2.text = b_desc
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = TEXT_BODY

# =============================================================
# SLIDE 8: LIVE DEMO TRANSITION SLIDE (WITH LIVE SCREENSHOT)
# =============================================================
s8 = prs.slides.add_slide(blank_slide_layout)
apply_lively_background(s8)
add_lively_header(s8, "System Demonstration: End-to-End Live Execution", category_pill="LIVE DEMO", pill_color=COLOR_EMERALD, pill_bg=COLOR_EMERALD_BG)

# Left Column: Step-by-Step Checklist
demo_card = create_lively_card(s8, Inches(0.8), Inches(1.5), Inches(5.8), Inches(5.4), title="LOCAL VERIFICATION CHECKLIST (/home/rishav/job-scout/)", accent_color=COLOR_EMERALD)
tb_demo = s8.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.4), Inches(4.7))
tf_demo = tb_demo.text_frame
tf_demo.word_wrap = True

demo_steps = [
    ("01. Container Health Check", "FastAPI (8001), Next.js (3000), PostgreSQL (5435), Redis (6379), Celery."),
    ("02. Ingest Resume in Vault", "Upload sample PDF resume; observe instant Markdown parsing & 384-d vector."),
    ("03. Live Matchmaking Query", "Execute live hybrid search; verify sub-10ms response via pgvector HNSW."),
    ("04. 100-Point ATS Diagnostic", "Inspect Radar module scorecard; review matched skills vs missing keywords."),
    ("05. Real-Time RAG Tailoring", "Trigger local Llama-3.2 tailoring; observe live SSE token streaming in browser."),
    ("06. 2.00 ms Instant Cache Hit", "Re-request same tailoring; verify instant 2.00 ms response from Redis cache."),
    ("07. Kanban Pipeline Transition", "Move job card to 'Applied'; verify persistent state change in PostgreSQL.")
]
for i, (s_title, s_desc) in enumerate(demo_steps):
    p = tf_demo.paragraphs[0] if i == 0 else tf_demo.add_paragraph()
    p.space_after = Pt(5)
    r1 = p.add_run()
    r1.text = s_title + ": "
    r1.font.bold = True
    r1.font.size = Pt(10)
    r1.font.color.rgb = COLOR_EMERALD
    r2 = p.add_run()
    r2.text = s_desc
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = TEXT_BODY

# Right Column: Embedded Screenshot of Live Copilot Module
ui_demo_card = create_lively_card(s8, Inches(6.9), Inches(1.5), Inches(5.6), Inches(5.4), title="LIVE SYSTEM INTERFACE (COPILOT MODULE)", accent_color=COLOR_INDIGO)
img_copilot = "/home/rishav/job-scout/assets/copilot.png"
if os.path.exists(img_copilot):
    s8.shapes.add_picture(img_copilot, Inches(7.1), Inches(2.1), Inches(5.2), Inches(4.5))

# =============================================================
# SLIDE 9: CONCLUSION & FUTURE HORIZONS
# =============================================================
s9 = prs.slides.add_slide(blank_slide_layout)
apply_lively_background(s9)
add_lively_header(s9, "Conclusion & Future Research Horizons", category_pill="VISION & ROADMAP", pill_color=COLOR_INDIGO, pill_bg=COLOR_INDIGO_BG)

create_lively_card(s9, Inches(0.8), Inches(1.5), Inches(5.7), Inches(5.4), title="PROJECT ACHIEVEMENTS & CONTRIBUTIONS", accent_color=COLOR_INDIGO)
tb_con = s9.shapes.add_textbox(Inches(1.0), Inches(2.1), Inches(5.3), Inches(4.6))
tf_con = tb_con.text_frame
tf_con.word_wrap = True

achievements = [
    ("Applicant-Side ATS Parity", "Democratized commercial-grade ATS screening intelligence, enabling candidates to navigate automated corporate filters equitably."),
    ("Sub-10ms Hybrid Retrieval at Scale", "Proved that standard PostgreSQL instances can deliver enterprise-grade 9.08 ms vector search across 100k rows (98.1x speedup)."),
    ("Uncompromising Data Sovereignty", "Demonstrated that cutting-edge AI recruitment tools do not require third-party cloud APIs, recurring SaaS paywalls, or PII leakage."),
    ("Rigorously Validated Engineering", "Validated through 126 automated tests (79% coverage), 95.65% HNSW recall, and blind human evaluation ratings of 4.20/5.00.")
]
for i, (head, desc) in enumerate(achievements):
    p = tf_con.paragraphs[0] if i == 0 else tf_con.add_paragraph()
    p.space_after = Pt(12)
    r1 = p.add_run()
    r1.text = "🎯 " + head + "\n"
    r1.font.bold = True
    r1.font.size = Pt(11)
    r1.font.color.rgb = COLOR_INDIGO
    r2 = p.add_run()
    r2.text = desc
    r2.font.size = Pt(10)
    r2.font.color.rgb = TEXT_BODY

create_lively_card(s9, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.4), title="FUTURE RESEARCH & SCALING ROADMAP", accent_color=COLOR_PURPLE)
tb_fut = s9.shapes.add_textbox(Inches(7.0), Inches(2.1), Inches(5.3), Inches(4.6))
tf_fut = tb_fut.text_frame
tf_fut.word_wrap = True

futures = [
    ("4-Bit GGUF Quantization (llama.cpp)", "Compress model RAM footprint to 2.2 GB and double CPU token generation throughput from 1.8 to 4.5+ tokens/second."),
    ("Multi-Profile Portfolio Management", "Allow applicants to maintain targeted resume variants (e.g., DevOps vs Data Science) with dynamic variant matchmaking."),
    ("Automated Headless LaTeX Engine", "Programmatically compile tailored Markdown resumes into pixel-perfect, ATS-compliant PDFs via headless Typst or LaTeX."),
    ("LoRA Domain Adaptation", "Fine-tune low-rank adapters on curated regional tech job postings to further reduce unsupported claims below 5%.")
]
for i, (head, desc) in enumerate(futures):
    p = tf_fut.paragraphs[0] if i == 0 else tf_fut.add_paragraph()
    p.space_after = Pt(12)
    r1 = p.add_run()
    r1.text = "🚀 " + head + "\n"
    r1.font.bold = True
    r1.font.size = Pt(11)
    r1.font.color.rgb = COLOR_PURPLE
    r2 = p.add_run()
    r2.text = desc
    r2.font.size = Pt(10)
    r2.font.color.rgb = TEXT_BODY

# Save Lively Deck
OUT_PPTX = "/home/rishav/job-scout/Document/Presentation/Job_Scout_Defense_Presentation.pptx"
prs.save(OUT_PPTX)
print(f"Lively Presentation saved successfully to {OUT_PPTX}")

# Convert to PDF
cmd = f"libreoffice --headless --convert-to pdf {OUT_PPTX} --outdir /home/rishav/job-scout/Document/Presentation/"
res = subprocess.run(cmd, shell=True, capture_output=True, text=True)
if res.returncode == 0:
    print("Lively PDF version successfully generated!")
else:
    print("PDF export error:", res.stderr)
